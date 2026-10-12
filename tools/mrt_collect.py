"""마이리얼트립 파트너 API로 여행 콘텐츠 재료(상품 + 제휴 링크)를 수집한다.

사용법:
  python mrt_collect.py 후쿠오카 --overseas --checkin 2026-11-13 --nights 2
  python mrt_collect.py 경주 --domestic

결과: 여행콘텐츠/data/<지역>_<날짜>.json, .xlsx
API 키: 같은 폴더 .env 의 MRT_API_KEY (저장소·블로그에 절대 올리지 말 것)
"""
import argparse, datetime as dt, json, os, re, sys, time
from pathlib import Path
import requests

BASE = "https://partner-ext-api.myrealtrip.com"
HERE = Path(__file__).resolve().parent

# 투어·티켓 검색 대상 카테고리 (도시별 categories API 값 중 있는 것만 사용)
PRIORITY = ["tour", "ticket", "usimwifi", "activity", "snap_v2", "transportation", "transportation_v2", "foodie", "spamassage", "class"]
# 수수료 API(/v1/revenues/commission-rates) 2026-10-11 조회 기준 카테고리별 요율 — 참고용
RATE_HINT = {"투어": 6.0, "스냅촬영": 7.0, "액티비티": 4.0, "클래스": 2.5, "입장권": 2.0, "티켓·입장권": 2.0, "숙소": 2.0, "이동·교통": 1.0, "이동/교통편의": 1.0}


def load_key():
    key = os.environ.get("MRT_API_KEY")
    env = HERE / ".env"
    if not key and env.exists():
        for line in env.read_text(encoding="utf-8").splitlines():
            if line.startswith("MRT_API_KEY="):
                key = line.split("=", 1)[1].strip()
    if not key:
        sys.exit("MRT_API_KEY가 없습니다 (.env 확인)")
    return key


class MRT:
    def __init__(self, key):
        self.s = requests.Session()
        self.s.headers.update({"Authorization": f"Bearer {key}", "Content-Type": "application/json"})
        self.link_cache = {}

    def post(self, path, body, pause=0.4):
        for attempt in range(3):
            r = self.s.post(BASE + path, json=body, timeout=30)
            if r.status_code == 429:
                time.sleep(15); continue
            time.sleep(pause)
            j = r.json()
            if r.status_code >= 400:
                raise RuntimeError(f"{path} {r.status_code} {j.get('result')}")
            return j["data"]
        raise RuntimeError(f"{path} rate limited")

    def categories(self, city):
        try:
            return self.post("/v1/products/tna/categories", {"city": city})["categories"]
        except Exception:
            return []

    def tna_search(self, keyword, category=None, size=5, sort="selling_count_desc"):
        body = {"keyword": keyword, "sort": sort, "page": 1, "size": size}
        if category: body["category"] = category
        return self.post("/v1/products/tna/search", body).get("items", [])

    def tna_detail(self, gid, date):
        try:
            return self.post("/v1/products/tna/detail", {"gid": gid, "selectedDate": date})
        except Exception:
            return {}

    def region(self, keyword, domestic):
        regs = self.post("/v1/products/accommodation/region-autocomplete", {"keyword": keyword, "isDomestic": domestic}, pause=1.3).get("regions", [])
        city = [r for r in regs if r.get("type") == "CITY"] or regs
        return city[0] if city else None

    def stays(self, region_id, checkin, checkout, size=10):
        return self.post("/v1/products/accommodation/search", {
            "regionId": region_id, "checkIn": checkin, "checkOut": checkout, "adultCount": 2, "page": 0, "size": size}, pause=1.3).get("items", [])

    def mylink(self, url):
        if url in self.link_cache: return self.link_cache[url]
        d = self.post("/v1/mylink", {"targetUrl": url})
        self.link_cache[url] = d["mylink"]
        return d["mylink"]


def strip_html(t):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", t or "")).strip()


def collect(keyword, domestic, checkin, nights, per_cat):
    api = MRT(load_key())
    today = dt.date.today().isoformat()
    pack = {"keyword": keyword, "domestic": domestic, "collected_at": today,
            "price_note": f"가격·평점은 {today} 마이리얼트립 기준 (변동 가능)", "tna": [], "stays": []}

    cats = {c["value"]: c["name"] for c in api.categories(keyword)}
    targets = [c for c in PRIORITY if c in cats] or [None]
    seen = set()
    for cat in targets:
        kept = 0
        for it in api.tna_search(keyword, cat, size=per_cat * 2):
            if kept >= per_cat: break
            if it["gid"] in seen: continue
            # 다른 도시 상품 제외: 설명("도시 ∙ 카테고리") 또는 상품명에 키워드가 있어야 함
            if keyword not in (it.get("description") or "") and keyword not in it["itemName"]: continue
            seen.add(it["gid"])
            det = api.tna_detail(it["gid"], checkin)
            cname = cats.get(cat, it.get("category", ""))
            pack["tna"].append({
                "group": cname, "gid": it["gid"], "name": it["itemName"], "category": it.get("category"),
                "price": it.get("salePrice"), "price_display": it.get("priceDisplay"),
                "review_score": round(it.get("reviewScore") or 0, 1), "review_count": it.get("reviewCount"),
                "included": det.get("included", []), "excluded": det.get("excluded", []),
                "itinerary": [f"{x.get('title','')}: {strip_html(x.get('description',''))}" for x in det.get("itineraries", [])],
                "tags": it.get("tags", []), "image": it.get("imageUrl"), "url": it.get("productUrl"),
                "affiliate_link": api.mylink(it["productUrl"]),
                "rate_hint": RATE_HINT.get(cname) or RATE_HINT.get(it.get("category", "")),
            })
            kept += 1

    reg = api.region(keyword, domestic)
    if reg:
        co = (dt.date.fromisoformat(checkin) + dt.timedelta(days=nights)).isoformat()
        pack["stay_condition"] = f"{reg['subName']} · {checkin}~{co} · 성인 2"
        for s in api.stays(reg["regionId"], checkin, co):
            if not s.get("salePrice"): continue  # 판매가 없음(마감 등) 제외
            pack["stays"].append({
                "name": s["itemName"], "star": s.get("starRating"), "price": s.get("salePrice"),
                "original_price": s.get("originalPrice"), "review_score": s.get("reviewScore"),
                "review_count": s.get("reviewCount"), "image": s.get("imageUrl"), "url": s.get("productUrl"),
                "affiliate_link": api.mylink(s["productUrl"]) if s.get("productUrl") else None,
            })
    return pack


def save(pack, outdir):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    outdir.mkdir(parents=True, exist_ok=True)
    stem = f"{pack['keyword']}_{pack['collected_at']}"
    (outdir / f"{stem}.json").write_text(json.dumps(pack, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = Workbook(); F = "맑은 고딕"
    hf = Font(name=F, bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="2F3A4A")
    def sheet(ws, title, headers, rows, widths):
        ws.title = title
        ws.append([pack["price_note"] + (" · " + pack.get("stay_condition", "") if title == "숙소" else "")])
        ws["A1"].font = Font(name=F, italic=True, color="666666", size=9)
        ws.append(headers)
        for c in ws[2]: c.font = hf; c.fill = fill; c.alignment = Alignment(horizontal="center")
        for r in rows: ws.append(r)
        for row in ws.iter_rows(min_row=3):
            for c in row: c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A3"
    sheet(wb.active, "투어·티켓", ["구분", "상품명", "시작가(원)", "평점", "리뷰수", "포함 사항", "수수료 참고(%)", "제휴 링크", "이미지"],
          [[t["group"], t["name"], t["price"], t["review_score"], t["review_count"], " / ".join(t["included"]), t["rate_hint"], t["affiliate_link"], t["image"]] for t in pack["tna"]],
          [14, 46, 12, 7, 9, 40, 12, 28, 30])
    sheet(wb.create_sheet(), "숙소", ["숙소명", "성급", "판매가(원, 조회 기간 기준)", "정가(원)", "평점", "리뷰수", "제휴 링크", "이미지"],
          [[s["name"], s["star"], s["price"], s["original_price"], s["review_score"], s["review_count"], s["affiliate_link"], s["image"]] for s in pack["stays"]],
          [40, 7, 14, 12, 7, 9, 28, 30])
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=3):
            for c in row:
                if isinstance(c.value, int) and c.column_letter in "CD": c.number_format = "#,##0"
    wb.save(outdir / f"{stem}.xlsx")
    return outdir / f"{stem}.json", outdir / f"{stem}.xlsx"


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("keyword")
    g = ap.add_mutually_exclusive_group(); g.add_argument("--domestic", action="store_true"); g.add_argument("--overseas", action="store_true")
    ap.add_argument("--checkin", default=(dt.date.today() + dt.timedelta(days=30)).isoformat())
    ap.add_argument("--nights", type=int, default=2)
    ap.add_argument("--per-cat", type=int, default=4)
    ap.add_argument("--out", default=str(HERE / "data"))
    a = ap.parse_args()
    pack = collect(a.keyword, not a.overseas, a.checkin, a.nights, a.per_cat)
    j, x = save(pack, Path(a.out))
    print(f"투어·티켓 {len(pack['tna'])}개, 숙소 {len(pack['stays'])}개 → {x}")
